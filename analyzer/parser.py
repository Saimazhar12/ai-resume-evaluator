"""
Resume text extraction.

Supports PDF (via pypdf) and DOCX (via python-docx). Takes a Werkzeug
FileStorage (as received from a Flask request) and returns plain text.
"""
from __future__ import annotations

import io

from pypdf import PdfReader
from docx import Document

ALLOWED_EXTENSIONS = {"pdf", "docx"}


class UnsupportedFileType(ValueError):
    pass


class EmptyResumeError(ValueError):
    pass


def _extension(filename: str) -> str:
    return filename.rsplit(".", 1)[-1].lower() if "." in filename else ""


def allowed_file(filename: str) -> bool:
    return _extension(filename) in ALLOWED_EXTENSIONS


def _extract_pdf(file_bytes: bytes) -> str:
    reader = PdfReader(io.BytesIO(file_bytes))
    pages = []
    for page in reader.pages:
        pages.append(page.extract_text() or "")
    return "\n".join(pages)


def _extract_docx(file_bytes: bytes) -> str:
    document = Document(io.BytesIO(file_bytes))

    parts = [p.text for p in document.paragraphs if p.text.strip()]

    # Tables often hold skills/experience in resume templates - include them.
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    return "\n".join(parts)


def extract_text(file_storage) -> str:
    """
    file_storage: a Flask/Werkzeug FileStorage object (request.files["resume"])
    Returns extracted plain text, raises UnsupportedFileType / EmptyResumeError.
    """
    filename = file_storage.filename or ""
    ext = _extension(filename)

    if ext not in ALLOWED_EXTENSIONS:
        raise UnsupportedFileType(
            f"Unsupported file type '.{ext}'. Please upload a PDF or DOCX resume."
        )

    file_bytes = file_storage.read()
    file_storage.seek(0)  # reset pointer in case caller needs it again

    if not file_bytes:
        raise EmptyResumeError("The uploaded file is empty.")

    if ext == "pdf":
        text = _extract_pdf(file_bytes)
    else:
        text = _extract_docx(file_bytes)

    text = text.strip()
    if len(text) < 30:
        raise EmptyResumeError(
            "Couldn't extract readable text from this file. "
            "If it's a scanned/image-based PDF, try a text-based export instead."
        )

    return text
